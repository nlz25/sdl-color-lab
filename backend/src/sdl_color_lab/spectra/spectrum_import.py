"""Read original reader workbooks without modifying or silently cleaning them.

Mappings are JSON serializable and use zero-based, inclusive cell coordinates.
``wells_rows``: wavelengths are on ``header_row``; ``axis_index`` is the
well-name column. ``wells_columns``: well names are on ``header_row``;
``axis_index`` is the wavelength column. Only the requested well is extracted.

The adapter does not subtract blanks, infer missing observations, or assert
that a file came from a particular instrument. Discovery is a convenience;
extraction validates the mapping against the actual uploaded bytes every time.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
from datetime import date, datetime
from pathlib import Path

from sdl_color_lab.spectra.spectrum_to_rgb import absorbance_to_color, colour_band_indices

CSV_SHEET = "__csv__"
ORIENTATIONS = ("wells_rows", "wells_columns")
_FIELDS = (
    "sheet",
    "orientation",
    "header_row",
    "first_data_row",
    "last_data_row",
    "first_data_col",
    "last_data_col",
    "axis_index",
)
_CSV_HEADERS = {
    ("wavelength_nm", "absorbance"),
    ("Wavelength(nm)", "OD Value"),
}


def normalize_well(value: object) -> str:
    """Normalize an explicit 96-well address, including e.g. ``A01``."""
    match = re.fullmatch(r"\s*([A-Ha-h])\s*0?([1-9]|1[0-2])\s*", str(value))
    if not match:
        raise ValueError(f"孔号无效：{value!r}；请使用 A1–H12。")
    return f"{match.group(1).upper()}{int(match.group(2))}"


def _well(value: object) -> str | None:
    try:
        return normalize_well(value)
    except ValueError:
        return None


def _json_cell(value: object) -> object:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def load_workbook_data(raw: bytes, filename: str) -> dict[str, list[list]]:
    """Read an immutable file snapshot; original bytes remain owned by caller."""
    if not isinstance(raw, bytes) or not raw:
        raise ValueError("上传文件为空，或不是字节数据。")
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".csv":
            text = raw.decode("utf-8-sig")
            header = next((line for line in text.splitlines() if line.strip()), "")
            # Tecan's European export separates columns with semicolons and uses
            # decimal commas. Do not let a decimal comma split an OD value.
            delimiter = ";" if ";" in header else "\t" if "\t" in header else ","
            rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
            while rows and not any(cell.strip() for cell in rows[0]):
                rows.pop(0)
            return {CSV_SHEET: rows}
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            # No formula evaluation; missing cached formula values must be
            # reported as missing data during extraction, not substituted.
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            try:
                return {
                    sheet.title: [
                        [_json_cell(cell) for cell in row]
                        for row in sheet.iter_rows(values_only=True)
                    ]
                    for sheet in workbook.worksheets
                }
            finally:
                workbook.close()
        if suffix == ".xls":
            import xlrd

            workbook = xlrd.open_workbook(file_contents=raw, on_demand=True)
            try:
                sheets = {}
                for sheet in workbook.sheets():
                    rows = []
                    for row_index in range(sheet.nrows):
                        row = []
                        for cell in sheet.row(row_index):
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                value = xlrd.xldate_as_datetime(
                                    cell.value, workbook.datemode
                                ).isoformat()
                            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                                value = bool(cell.value)
                            elif cell.ctype == xlrd.XL_CELL_ERROR:
                                value = xlrd.error_text_from_code.get(
                                    cell.value, "#EXCEL_ERROR"
                                )
                            else:
                                value = _json_cell(cell.value)
                            row.append(value)
                        rows.append(row)
                    sheets[sheet.name] = rows
                return sheets
            finally:
                workbook.release_resources()
    except (ImportError, ModuleNotFoundError):
        raise
    except Exception as exc:
        raise ValueError(f"无法读取 {filename}：{exc}") from exc
    raise ValueError("仅支持 .xlsx、.xls 或 UTF-8 编码的 .csv 文件。")


def _cell(rows: list[list], row: int, column: int) -> object:
    if 0 <= row < len(rows) and 0 <= column < len(rows[row]):
        return rows[row][column]
    return None


def _is_empty(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _number(value: object, description: str) -> float:
    if isinstance(value, bool) or _is_empty(value):
        raise ValueError(f"{description}缺失或不是有效数字。")
    if isinstance(value, str):
        value = value.strip()
        if re.fullmatch(r"[+-]?(?:\d+,\d+|,\d+)(?:[eE][+-]?\d+)?", value):
            value = value.replace(",", ".")
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(
            f"{description}不是有效数字：{value!r}；不能忽略缺失值、溢出或饱和标记。"
        ) from exc
    if not math.isfinite(result):
        raise ValueError(f"{description}必须是有限数值，不能包含 NaN 或无穷大。")
    return result


def _wavelength(value: object, description: str = "波长") -> float:
    # Tecan/hand-mapped headers may explicitly include their unit.
    if isinstance(value, str):
        match = re.fullmatch(r"\s*(\d+(?:\.\d+)?)\s*(?:nm)?\s*", value, re.IGNORECASE)
        if match:
            value = match.group(1)
    return _number(value, description)


def _looks_like_wavelength(value: object) -> bool:
    try:
        return 200 <= _wavelength(value) <= 1100
    except ValueError:
        return False


def _hash_json(value: object) -> str:
    return hashlib.sha256(
        json.dumps(
            value,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            allow_nan=False,
        ).encode("utf-8")
    ).hexdigest()


def create_block(
    sheets: dict[str, list[list]],
    sheet: str,
    orientation: str,
    header_row: int,
    first_data_row: int,
    last_data_row: int,
    first_data_col: int,
    last_data_col: int,
    axis_index: int,
) -> dict:
    """Create a reusable mapping from the selected rectangle and explicit axes.

    Data bounds exclude header cells and the axis column. In wells_rows they
    enclose absorbance values across well rows and wavelength columns; in
    wells_columns they enclose absorbance across wavelength rows and well
    columns. Invalid measurements are surfaced by extract_spectrum.
    """
    if sheet not in sheets:
        raise ValueError(f"工作表不存在：{sheet}")
    if orientation not in ORIENTATIONS:
        raise ValueError("请选择孔按行或孔按列的光谱排列方向。")
    indices = (
        header_row,
        first_data_row,
        last_data_row,
        first_data_col,
        last_data_col,
        axis_index,
    )
    if any(isinstance(x, bool) or not isinstance(x, int) or x < 0 for x in indices):
        raise ValueError("映射坐标必须是从 0 开始的非负整数。")
    rows = sheets[sheet]
    width = max((len(row) for row in rows), default=0)
    if (
        first_data_row > last_data_row
        or first_data_col > last_data_col
        or last_data_row >= len(rows)
        or header_row >= first_data_row
        or last_data_col >= width
        or axis_index >= width
        or first_data_col <= axis_index <= last_data_col
    ):
        raise ValueError("映射区域超出表格、为空，或与孔号／波长轴重叠。")
    mapping = dict(zip(_FIELDS, (sheet, orientation, *indices)))
    # zip above intentionally follows the public argument order.
    positions = (
        (
            _cell(rows, row, axis_index)
            for row in range(first_data_row, last_data_row + 1)
        )
        if orientation == "wells_rows"
        else (
            _cell(rows, header_row, column)
            for column in range(first_data_col, last_data_col + 1)
        )
    )
    wells = [well for value in positions if (well := _well(value)) is not None]
    if not wells:
        raise ValueError("所选区域未找到明确的孔号（例如 A1 或 A01）。")
    block_id = _hash_json(mapping)[:16]
    return {
        **mapping,
        "id": block_id,
        "block_id": block_id,
        "wells": list(dict.fromkeys(wells)),
        "label": (
            f"{sheet} · 第 {header_row + 1} 行表头 · "
            f"{'孔按行' if orientation == 'wells_rows' else '孔按列'}"
        ),
    }


def _next_data_end(rows: list[list], start: int, axis: int, kind: str) -> int:
    """Find a contiguous table body without dropping interior missing cells."""
    last = start - 1
    for index in range(start, len(rows)):
        row = rows[index]
        if all(_is_empty(value) for value in row):
            break
        value = _cell(rows, index, axis)
        valid = (
            _well(value) is not None
            if kind == "well"
            else _looks_like_wavelength(value)
        )
        if not valid:
            # A data row with a blank/malformed wavelength must remain in the
            # rectangle so validation reports it. A new labeled header or
            # metadata line terminates a scan.
            if kind == "wave" and (
                _is_empty(value) or _numeric_but_not_wavelength(value)
            ):
                last = index
                continue
            break
        last = index
    return last


def _numeric_but_not_wavelength(value: object) -> bool:
    try:
        _number(value, "波长")
        return True
    except ValueError:
        return False


def discover_blocks(sheets: dict[str, list[list]]) -> list[dict]:
    """Discover explicit well/wavelength tables, never choose a scan for caller."""
    if set(sheets) == {CSV_SHEET}:
        rows = sheets[CSV_SHEET]
        if not rows:
            raise ValueError("上传文件为空")
        if tuple(cell.strip() for cell in rows[0]) in _CSV_HEADERS:
            return [
                {
                    "sheet": CSV_SHEET,
                    "orientation": "csv",
                    "id": "csv",
                    "block_id": "csv",
                    "wells": [],
                    "label": "单孔 CSV（需绑定孔号）",
                }
            ]
    blocks = []
    for sheet_name, rows in sheets.items():
        for header, row in enumerate(rows):
            # Wavelengths along columns. Include the full span, even an invalid
            # interior header, so extraction cannot quietly omit one point.
            wave_columns = [
                index
                for index, value in enumerate(row)
                if _looks_like_wavelength(value)
            ]
            if len(wave_columns) >= 2:
                first, last = min(wave_columns), max(wave_columns)
                for axis in range(first):
                    start = header + 1
                    if _well(_cell(rows, start, axis)) is None:
                        continue
                    end = _next_data_end(rows, start, axis, "well")
                    blocks.append(
                        create_block(
                            sheets,
                            sheet_name,
                            "wells_rows",
                            header,
                            start,
                            end,
                            first,
                            last,
                            axis,
                        )
                    )
            # Well names along columns. Keep invalid interior names within the
            # rectangle; an explicitly selected well is still matched by name.
            well_columns = [index for index, value in enumerate(row) if _well(value)]
            if well_columns:
                first, last = min(well_columns), max(well_columns)
                for axis in range(first):
                    start = header + 1
                    if not _looks_like_wavelength(_cell(rows, start, axis)):
                        continue
                    end = _next_data_end(rows, start, axis, "wave")
                    if end - start >= 1:
                        blocks.append(
                            create_block(
                                sheets,
                                sheet_name,
                                "wells_columns",
                                header,
                                start,
                                end,
                                first,
                                last,
                                axis,
                            )
                        )
    # Multiple axes may represent distinct tables. Preserve their identity and
    # require the user to resolve that ambiguity rather than picking the first.
    discovered = list({block["id"]: block for block in blocks}.values())
    if not discovered and set(sheets) == {CSV_SHEET}:
        raise ValueError("CSV 须包含波长、吸光度两列，或带明确孔号与波长表头的 Tecan 多孔表格。")
    return discovered


def _timestamp(rows: list[list], block: dict) -> str | None:
    """Return only explicitly timezone-bearing, scan-local measurement times."""
    header = block.get("header_row", 0)
    lower = max(0, header - 12)
    # Do not carry timestamps across another scan's explicit data/header rows.
    for index in range(header - 1, lower - 1, -1):
        row = rows[index]
        if (
            any(_well(cell) for cell in row)
            or sum(_looks_like_wavelength(cell) for cell in row) >= 2
        ):
            lower = index + 1
            break
    candidates = []
    for row in rows[lower:header]:
        label = " ".join(str(value) for value in row if value is not None)
        # An export/save timestamp is not the acquisition time. Likewise a
        # generic workbook "Date" field is insufficient to establish provenance.
        if re.search(
            r"export|import|download|saved|created|modified|print|导出|保存",
            label,
            re.I,
        ):
            continue
        if not re.search(
            r"measurement|started|measured|acquired|scan time|测量|采集|检测",
            label,
            re.I,
        ):
            continue
        for value in row:
            if not isinstance(value, str):
                continue
            # Require ISO date/time with Z or ±HH:MM. A naive Excel datetime
            # is retained in preview, but must be confirmed by an operator.
            for match in re.finditer(
                r"\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(?::\d{2}(?:\.\d+)?)?"
                r"(?:Z|[+-]\d{2}:\d{2})",
                value,
            ):
                try:
                    parsed = datetime.fromisoformat(
                        match.group(0).replace("Z", "+00:00")
                    )
                except ValueError:
                    continue
                if parsed.tzinfo is not None and parsed.utcoffset() is not None:
                    candidates.append(parsed.isoformat())
    unique = list(dict.fromkeys(candidates))
    return unique[0] if len(unique) == 1 else None


def _validated_color(wavelengths: list[float], absorbance: list[float]) -> dict:
    if len(wavelengths) < 2 or len(wavelengths) != len(absorbance):
        raise ValueError("光谱至少需要两个波长点，波长与吸光度数量必须一致。")
    gaps = [right - left for left, right in zip(wavelengths, wavelengths[1:])]
    if any(gap <= 0 for gap in gaps):
        raise ValueError("波长必须严格递增且不能重复；请检查扫描区域和排列方向。")
    band = colour_band_indices(wavelengths)
    if wavelengths[0] > 380 or wavelengths[-1] < 780:
        raise ValueError("光谱必须完整覆盖 380–780 nm。")
    if any(absorbance[index] < 0 for index in band):
        raise ValueError("吸光度包含负值，请核对仪器的纯水／孔板背景处理。")
    return absorbance_to_color(wavelengths, absorbance)


def extract_spectrum(
    raw: bytes,
    filename: str,
    block: dict,
    well: str,
    expected_well: str | None = None,
    *,
    require_color: bool = True,
) -> dict:
    """Extract, validate and convert exactly one explicitly identified well."""
    requested_well = normalize_well(well)
    if expected_well is not None and requested_well != normalize_well(expected_well):
        raise ValueError(f"反馈孔号必须为当前待测孔 {normalize_well(expected_well)}。")
    sheets = load_workbook_data(raw, filename)
    orientation = block.get("orientation")
    acquired_at = None
    source_well = requested_well
    if orientation == "csv":
        if set(sheets) != {CSV_SHEET}:
            raise ValueError("CSV 映射不能用于 Excel 文件。")
        canonical = discover_blocks(sheets)[0]
        rows = sheets[CSV_SHEET]
        wavelengths, values = [], []
        for index, row in enumerate(rows[1:], start=2):
            if len(row) != 2:
                raise ValueError(f"CSV 第 {index} 行必须恰好包含波长和吸光度两列。")
            wavelengths.append(_wavelength(row[0], f"CSV 第 {index} 行波长"))
            values.append(row[1])
        source_well = None  # A two-column file cannot independently attest a well.
        mapping = {"sheet": CSV_SHEET, "orientation": "csv"}
    else:
        try:
            canonical = create_block(sheets, **{key: block[key] for key in _FIELDS})
        except KeyError as exc:
            raise ValueError(f"光谱映射缺少字段：{exc.args[0]}") from exc
        mapping = {key: canonical[key] for key in _FIELDS}
        rows = sheets[canonical["sheet"]]
        start, end = canonical["first_data_row"], canonical["last_data_row"]
        first, last = canonical["first_data_col"], canonical["last_data_col"]
        header, axis = canonical["header_row"], canonical["axis_index"]
        if orientation == "wells_rows":
            matches = [
                index
                for index in range(start, end + 1)
                if _well(_cell(rows, index, axis)) == requested_well
            ]
            if len(matches) != 1:
                raise ValueError(
                    f"所选扫描区域须恰好包含一次孔 {requested_well}，实际为 {len(matches)}。"
                )
            selected_row = matches[0]
            wavelengths = [
                _wavelength(
                    _cell(rows, header, col), f"第 {header + 1} 行第 {col + 1} 列波长"
                )
                for col in range(first, last + 1)
            ]
            values = [_cell(rows, selected_row, col) for col in range(first, last + 1)]
        else:
            matches = [
                index
                for index in range(first, last + 1)
                if _well(_cell(rows, header, index)) == requested_well
            ]
            if len(matches) != 1:
                raise ValueError(
                    f"所选扫描区域须恰好包含一次孔 {requested_well}，实际为 {len(matches)}。"
                )
            selected_col = matches[0]
            wavelengths = [
                _wavelength(_cell(rows, row, axis), f"第 {row + 1} 行波长")
                for row in range(start, end + 1)
            ]
            values = [_cell(rows, row, selected_col) for row in range(start, end + 1)]
        acquired_at = _timestamp(rows, canonical)
    if len(wavelengths) < 2:
        raise ValueError("光谱至少需要两个波长点。")
    gaps = [right - left for left, right in zip(wavelengths, wavelengths[1:])]
    if any(gap <= 0 for gap in gaps):
        raise ValueError("波长必须严格递增且不能重复。")
    original_range = [wavelengths[0], wavelengths[-1]]
    original_count = len(wavelengths)
    needed = set(colour_band_indices(wavelengths).tolist())
    valid_waves, absorbance, ignored = [], [], []
    for index, (wave, value) in enumerate(zip(wavelengths, values)):
        try:
            number = _number(value, f"孔 {requested_well} 波长 {wave:g} nm 的吸光度")
        except ValueError:
            # An explicit OVER outside the integration band is irrelevant to the
            # colour calculation. Preserve it in provenance; never replace it.
            if index not in needed and isinstance(value, str) and value.strip().upper() == "OVER":
                ignored.append({"wavelength_nm": wave, "value": value})
                continue
            raise
        valid_waves.append(wave)
        absorbance.append(number)
    wavelengths = valid_waves
    conversion_error = None
    try:
        color = _validated_color(wavelengths, absorbance)
    except ValueError as exc:
        if require_color:
            raise
        color, conversion_error = None, str(exc)
    instrument = next(
        (value.split(":", 1)[1].strip()
         for row in sheets[canonical["sheet"]][:28]
         for value in row
         if isinstance(value, str) and value.strip().lower().startswith("device:")),
        None,
    )
    return {
        "well": requested_well,
        "source_well": source_well,
        "filename": Path(filename).name,
        "wavelength_nm": wavelengths,
        "absorbance": absorbance,
        "color": color,
        "ready_for_rgb": color is not None,
        "conversion_error": conversion_error,
        "instrument": instrument,
        "original_wavelength_range_nm": original_range,
        "original_sample_count": original_count,
        "ignored_outside_colour_band": ignored,
        "source_sha256": hashlib.sha256(raw).hexdigest(),
        "spectrum_sha256": _hash_json(
            {
                "wavelength_nm": wavelengths,
                "absorbance": absorbance,
            }
        ),
        "sheet": canonical["sheet"],
        "block_id": canonical["block_id"],
        "mapping": mapping,
        "acquired_at": acquired_at,
        "sample_count": len(wavelengths),
        "wavelength_range_nm": [wavelengths[0], wavelengths[-1]],
        "sampling_interval_nm": (
            float(format((original_range[1] - original_range[0]) / (original_count - 1), ".12g"))
            if all(math.isclose(gap, gaps[0], rel_tol=1e-8, abs_tol=1e-8) for gap in gaps)
            else None
        ),
        "max_sampling_interval_nm": float(format(max(gaps), ".12g")),
    }


def normalized_csv(spectrum: dict) -> str:
    """Write measured numeric points only, with invariant decimal formatting."""
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, lineterminator="\n")
    writer.writerow(("wavelength_nm", "absorbance"))
    for wave, absorbance in zip(spectrum["wavelength_nm"], spectrum["absorbance"]):
        writer.writerow((format(wave, ".17g"), format(absorbance, ".17g")))
    return stream.getvalue()
